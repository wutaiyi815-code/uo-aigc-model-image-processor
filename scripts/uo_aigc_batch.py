import argparse
import base64
import json
import os
import re
import shutil
import sys
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path
from typing import Any

import requests
from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont


API_BASE = "https://toapis.com/v1"
DEFAULT_IMAGE_MODEL = "gemini-3-pro-image-preview"
DEFAULT_VISION_MODEL = "gemini-3.1-flash-lite-preview-official"
CLASS_JSON_NAME = "AIGC视觉识别.json"
CLASS_JSON_ASCII_NAME = "aigc_image_analysis.json"
SIZE = "3:4"
RESOLUTION = "2K"
REQUEST_TIMEOUT = 600
MAX_RETRIES = 3
POLL_INTERVAL = 8
IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".webp"}
SKIP_TERMS = [
    "留白", "分辨率", "尺寸", "加白底", "导出jpg", "jpg格式",
    "裁剪", "修剪", "只展示上半身", "只展示上身", "补齐图片后", "修改图片后",
]

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")


def now() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def sanitize_name(text: str) -> str:
    text = re.sub(r'[<>:"/\\|?*\r\n]+', "_", text).strip(" ._")
    return text[:80] or "AIGC任务"


def request_with_retry(method: str, url: str, **kwargs) -> requests.Response:
    last_exc = None
    for attempt in range(1, MAX_RETRIES + 1):
        try:
            resp = requests.request(method, url, timeout=REQUEST_TIMEOUT, **kwargs)
            if resp.status_code < 500:
                return resp
            last_exc = RuntimeError(f"HTTP {resp.status_code}: {resp.text[:500]}")
        except Exception as exc:
            last_exc = exc
        if attempt < MAX_RETRIES:
            time.sleep(2 * attempt)
    raise RuntimeError(f"request failed after {MAX_RETRIES} attempts: {last_exc}")


def safe_json(resp: requests.Response) -> Any:
    try:
        return resp.json()
    except Exception:
        return {"text": resp.text}


def append_log(path: Path, record: dict[str, Any]) -> None:
    with path.open("a", encoding="utf-8") as fh:
        fh.write("\n" + "=" * 80 + "\n")
        fh.write(json.dumps(record, ensure_ascii=False, indent=2, default=str))
        fh.write("\n")


def modification_items(text: str) -> list[str]:
    value = str(text or "").replace("&#x20;", " ").strip()
    if not value:
        return []
    pattern = re.compile(r"(?:^|\n)\s*\d+\s*[\.、）\)]\s*(.*?)(?=(?:\n\s*\d+\s*[\.、）\)])|\Z)", re.S)
    numbered = [" ".join(m.group(1).split()) for m in pattern.finditer(value) if m.group(1).strip()]
    if numbered:
        return numbered
    lines = [" ".join(line.split()) for line in value.splitlines() if line.strip()]
    return lines or [" ".join(value.split())]


def skip_reason(item: str) -> str | None:
    compact = re.sub(r"\s+", "", item).lower()
    for term in SKIP_TERMS:
        if term in compact:
            return f"excluded_non_generation_edit:{term}"
    if re.search(r"\d+\s*[\\*x×]\s*\d+", compact):
        return "excluded_resolution_or_dimensions"
    return None


def is_actionable_model_task(item: str) -> bool:
    missing_or_angle = "模特" in item and any(t in item for t in ["缺失", "缺少", "补充", "增加", "扩充"])
    pose = any(t in item for t in ["姿势", "动作"])
    recolor = "同款" in item and any(t in item for t in ["换色", "改色", "换"])
    garment_edit = "模特" in item and any(t in item for t in ["服装修改", "换装", "产品上身", "服装替换"])
    same_person = any(t in item for t in ["换成同一个人", "换成一个人", "替换为同一个模特"])
    return missing_or_angle or pose or recolor or garment_edit or same_person


def screen_modification_item(item: str) -> tuple[str | None, list[dict[str, str]]]:
    clauses = [part.strip() for part in re.split(r"[，,；;。]+", item) if part.strip()]
    actionable_parts = []
    skipped_parts = []
    for clause in clauses or [item]:
        reason = skip_reason(clause)
        if reason:
            skipped_parts.append({"item": clause, "reason": reason})
        elif is_actionable_model_task(clause):
            actionable_parts.append(clause)
        else:
            skipped_parts.append({"item": clause, "reason": "unsupported_task_type"})
    return ("，".join(actionable_parts) if actionable_parts else None), skipped_parts


def read_tasks(root: Path, excel_name: str) -> list[dict[str, Any]]:
    xlsx = root / excel_name
    if not xlsx.exists():
        candidates = list(root.glob("*.xlsx"))
        if len(candidates) == 1:
            xlsx = candidates[0]
        else:
            raise FileNotFoundError(f"Excel file not found: {xlsx}")
    wb = load_workbook(xlsx, data_only=True)
    ws = wb.active
    headers = [str(c.value or "").strip() for c in ws[1]]
    sku_col = next((i + 1 for i, h in enumerate(headers) if h.upper() == "SKU" or "SKU" in h.upper()), 2)
    mod_col = next((i + 1 for i, h in enumerate(headers) if "修改意见" in h), None) or 10
    gender_col = next((i + 1 for i, h in enumerate(headers) if "性别" in h), 5)
    color_col = next((i + 1 for i, h in enumerate(headers) if "颜色" in h or "色" == h), 6)
    tasks = []
    audit = []
    for row in range(2, ws.max_row + 1):
        sku = str(ws.cell(row, sku_col).value or "").strip()
        raw = ws.cell(row, mod_col).value
        if not sku or not raw:
            continue
        sku_dir = root / sku
        if not sku_dir.is_dir():
            continue
        items = modification_items(str(raw))
        actionable = []
        skipped = []
        for item in items:
            actionable_text, skipped_parts = screen_modification_item(item)
            if actionable_text:
                actionable.append(actionable_text)
            skipped.extend(skipped_parts)
        audit.append({"row": row, "sku": sku, "all_items": items, "actionable_items": actionable, "skipped_items": skipped})
        if not actionable:
            continue
        tasks.append(
            {
                "row": row,
                "sku": sku,
                "sku_dir": str(sku_dir),
                "gender": str(ws.cell(row, gender_col).value or "").strip(),
                "color": str(ws.cell(row, color_col).value or "").strip(),
                "raw_mod": str(raw),
                "first_mod": actionable[0],
                "actionable_mods": actionable,
                "skipped_mods": skipped,
            }
        )
    (root / "AIGC任务筛选.json").write_text(json.dumps({"created_at": now(), "rows": audit}, ensure_ascii=False, indent=2), encoding="utf-8")
    return tasks


def images_in_folder(sku_dir: Path) -> list[Path]:
    return [
        p
        for p in sorted(sku_dir.iterdir())
        if p.is_file()
        and p.suffix.lower() in IMAGE_EXTS
        and not p.name.startswith("AIGC+")
        and not p.name.lower().endswith(".txt")
    ]


def create_contact_sheet(root: Path, tasks: list[dict[str, Any]]) -> Path:
    out_dir = root / "AIGC视觉识别拼版"
    out_dir.mkdir(exist_ok=True)
    try:
        font = ImageFont.truetype("arial.ttf", 14)
    except Exception:
        font = ImageFont.load_default()
    for task in tasks:
        imgs = images_in_folder(Path(task["sku_dir"]))
        if not imgs:
            continue
        tile_w, tile_h, label_h, cols = 260, 340, 58, 4
        rows = (len(imgs) + cols - 1) // cols
        sheet = Image.new("RGB", (cols * tile_w, rows * (tile_h + label_h)), "white")
        draw = ImageDraw.Draw(sheet)
        for idx, img_path in enumerate(imgs):
            r, c = divmod(idx, cols)
            im = Image.open(img_path).convert("RGB")
            im.thumbnail((tile_w - 20, tile_h - 20), Image.LANCZOS)
            x = c * tile_w + (tile_w - im.width) // 2
            y = r * (tile_h + label_h) + 10
            sheet.paste(im, (x, y))
            draw.text((c * tile_w + 8, r * (tile_h + label_h) + tile_h + 4), f"{idx + 1}: {img_path.name}", fill="black", font=font)
            draw.text((c * tile_w + 8, r * (tile_h + label_h) + tile_h + 24), f"{img_path.stat().st_size // 1024}KB", fill="gray", font=font)
        sheet.save(out_dir / f"{task['sku']}.jpg", quality=92)
    return out_dir


def upload_image(path: Path, api_key: str) -> dict[str, Any]:
    with path.open("rb") as fh:
        resp = request_with_retry(
            "POST",
            f"{API_BASE}/uploads/images",
            headers={"Authorization": f"Bearer {api_key}"},
            files={"file": (path.name, fh)},
        )
    payload = safe_json(resp)
    if resp.status_code >= 400:
        raise RuntimeError(f"upload failed {resp.status_code}: {payload}")
    url = payload.get("data", {}).get("url") or payload.get("url")
    if not url:
        raise RuntimeError(f"upload response missing url: {payload}")
    return {"path": str(path), "url": url, "response": payload}


def classify_one_image(path: Path, api_key: str, vision_model: str) -> dict[str, Any]:
    upload = upload_image(path, api_key)
    prompt = (
        "Analyze this apparel image. Return compact JSON with keys: "
        "kind(model|flat|detail|other), view(front|side_3_4|back|unknown), "
        "garment_type(top|hoodie|pants|skirt|dress|set|unknown), model_gender(male|female|unknown), "
        "fit_notes(regular|fitted|boxy|oversized|baggy|jorts|loose|unknown plus short garment-specific notes), "
        "pose_notes, product_side(front|back|unknown), confidence."
    )
    body = {
        "model": vision_model,
        "messages": [
            {
                "role": "user",
                "content": [
                    {"type": "text", "text": prompt},
                    {"type": "image_url", "image_url": {"url": upload["url"]}},
                ],
            }
        ],
        "max_tokens": 300,
    }
    resp = request_with_retry(
        "POST",
        f"{API_BASE}/chat/completions",
        headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"},
        json=body,
    )
    payload = safe_json(resp)
    text = ""
    try:
        text = payload["choices"][0]["message"]["content"]
    except Exception:
        text = json.dumps(payload, ensure_ascii=False)
    parsed = {}
    m = re.search(r"\{.*\}", text, re.S)
    if m:
        try:
            parsed = json.loads(m.group(0))
        except Exception:
            parsed = {}
    parsed.update({"file": path.name, "path": str(path), "vision_raw": text, "upload_url": upload["url"]})
    return parsed


def classify(root: Path, tasks: list[dict[str, Any]], api_key: str | None, vision_model: str) -> Path:
    create_contact_sheet(root, tasks)
    output = root / CLASS_JSON_NAME
    data = {"created_at": now(), "skus": {}}
    for task in tasks:
        sku_dir = Path(task["sku_dir"])
        records = []
        for img in images_in_folder(sku_dir):
            if api_key:
                try:
                    records.append(classify_one_image(img, api_key, vision_model))
                except Exception as exc:
                    records.append({"file": img.name, "path": str(img), "kind": "unknown", "view": "unknown", "error": repr(exc)})
            else:
                records.append({"file": img.name, "path": str(img), "kind": "unknown", "view": "unknown", "fit_notes": "", "manual_note": "fill manually"})
        data["skus"][task["sku"]] = {
            "task": task["first_mod"],
            "actionable_tasks": task.get("actionable_mods", [task["first_mod"]]),
            "skipped_tasks": task.get("skipped_mods", []),
            "images": records,
            "front_flat": "",
            "back_flat": "",
            "base_model": "",
            "same_person_base": "",
            "same_person_targets": [],
            "eagle_base": "",
            "eagle_selection": {
                "path": "",
                "search_tags": [],
                "full_body": False,
                "background": "",
                "gender_match": False,
                "styling_compatible": False,
                "batch_unique": False,
                "uniqueness_signature": "",
                "pants_waistband_preference": "",
                "top_hem_can_hang_naturally": None,
                "manual_verified": False,
            },
            "custom_steps": [],
            "manual_verified": False,
        }
    output.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    (root / CLASS_JSON_ASCII_NAME).write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    return output


def load_classification(root: Path) -> dict[str, Any]:
    path = root / CLASS_JSON_NAME
    if not path.exists():
        path = root / CLASS_JSON_ASCII_NAME
    if not path.exists():
        raise FileNotFoundError(f"Run --mode classify and manually verify {CLASS_JSON_NAME} or {CLASS_JSON_ASCII_NAME} first")
    data = json.loads(path.read_text(encoding="utf-8-sig"))
    return data.get("skus", {})


def path_by_name(sku_dir: Path, name: str) -> Path:
    p = sku_dir / name
    if p.exists():
        return p
    for candidate in images_in_folder(sku_dir):
        if candidate.name == name:
            return candidate
    raise FileNotFoundError(f"reference image not found in {sku_dir}: {name}")


def missing_angles(first: str) -> list[str]:
    out = []
    if "正面" in first:
        out.append("front")
    if "侧面" in first:
        out.append("side")
    if "背面" in first:
        out.append("back")
    return out


def garment_word(class_info: dict[str, Any], fallback: str = "garment") -> str:
    for img in class_info.get("images", []):
        g = img.get("garment_type")
        if g and g != "unknown":
            return {"top": "top", "hoodie": "hoodie", "pants": "pants", "skirt": "skirt", "dress": "dress", "set": "outfit"}.get(g, fallback)
    return fallback


def fit_notes_for(class_info: dict[str, Any], angle: str | None = None) -> str:
    flat_name = None
    if angle == "back":
        flat_name = class_info.get("back_flat")
    elif angle in {"front", "side"}:
        flat_name = class_info.get("front_flat")
    candidates = []
    for img in class_info.get("images", []):
        if flat_name and img.get("file") == flat_name:
            candidates.insert(0, img)
        elif img.get("kind") == "flat":
            candidates.append(img)
    for img in candidates:
        note = str(img.get("fit_notes") or "").strip()
        if note and note.lower() != "unknown":
            return note
    return str(class_info.get("fit_notes") or "").strip()


def fit_clause(fit_notes: str) -> str:
    note = " ".join(str(fit_notes or "").split())
    if not note:
        return ""
    if " fit" in note.lower() or " silhouette" in note.lower():
        return f", preserving the Figure 2 {note}"
    return f", preserving the Figure 2 {note} fit"


def prompt_for(kind: str, garment: str, fit_notes: str = "", image_model: str = "") -> str:
    if kind == "angle_side":
        return "Adjust the model to a three-quarter side view in a full-body shot, with both arms hanging naturally by the sides and the head facing directly forward. Maintain the original head-to-body ratio and framing while ensuring the overall pose is natural and harmonious."
    if kind == "angle_back":
        if "gpt" in image_model.lower():
            return "生成背视图，人物重心放在其中一条腿上。"
        return "生成全身景别的背视图，人物重心放在其中一条腿上。"
    if kind == "angle_front":
        return "Adjust the character's pose to a full-body front view, with weight on one leg and a natural standing posture."
    if kind == "try_on":
        return f"Replace the {garment} worn by the person in Figure 1 with the {garment} from Figure 2{fit_clause(fit_notes)}, keeping the person's pose and all other clothing from Figure 1 unchanged."
    if kind == "color":
        return f"Change the color of the {garment} worn by the person in Figure 1 to match the {garment} in Figure 2{fit_clause(fit_notes)}, while keeping the silhouette, pattern placement, trims, and decorative details consistent with Figure 2."
    if kind == "pose":
        return "将图中人物动作更改为双手放下，自然垂于身体两侧，并调整为立正站姿。"
    if kind == "same_person":
        return "将图1的模特替换为图2的模特，并保持图1的动作不变。"
    if kind == "tattoo":
        return "去除图中模特手臂上的纹身，其他内容保持不变。"
    return "按修改意见对图片进行对应修改。"


def validate_generation_prompt(prompt: str, sku: str = "") -> None:
    text = str(prompt or "").strip()
    if not text:
        raise ValueError("Generation prompt is empty")
    forbidden = []
    if re.search(r"\bSKU\b", text, re.I) or (sku and sku.lower() in text.lower()):
        forbidden.append("SKU")
    if re.search(r"(?:[A-Za-z]:[\\/]|\\\\[^\s]+[\\/])", text):
        forbidden.append("file path")
    checks = {
        "image ratio": r"\b\d+\s*:\s*\d+\b",
        "resolution/clarity parameter": r"(?:分辨率|清晰度|像素|\b[1248]K\b|高清)",
        "API execution detail": r"(?:\bAPI\b|gpt-image|gemini|超时|重试|输出路径|保存到)",
    }
    for label, pattern in checks.items():
        if re.search(pattern, text, re.I):
            forbidden.append(label)
    if forbidden:
        raise ValueError(f"Prompt contains non-visual task information: {', '.join(forbidden)}")


def model_ref_for(class_info: dict[str, Any], sku_dir: Path, preferred: str | None = None) -> Path:
    if preferred:
        return path_by_name(sku_dir, preferred)
    if class_info.get("base_model"):
        return path_by_name(sku_dir, class_info["base_model"])
    for img in class_info.get("images", []):
        if img.get("kind") == "model":
            return path_by_name(sku_dir, img["file"])
    raise FileNotFoundError("No model reference found")


def model_ref_for_angle(class_info: dict[str, Any], sku_dir: Path, angle: str) -> Path | None:
    view_map = {
        "front": ["front"],
        "side": ["side_3_4", "side"],
        "back": ["back"],
    }
    for view in view_map.get(angle, []):
        for img in class_info.get("images", []):
            if img.get("kind") == "model" and img.get("view") == view:
                return path_by_name(sku_dir, img["file"])
    return None


def flat_for(class_info: dict[str, Any], sku_dir: Path, angle: str) -> Path:
    key = "back_flat" if angle == "back" else "front_flat"
    name = class_info.get(key) or class_info.get("front_flat")
    if not name:
        raise FileNotFoundError(f"Missing {key} in AIGC视觉识别.json")
    return path_by_name(sku_dir, name)


def back_angle_refs(class_info: dict[str, Any], sku_dir: Path) -> list[Path]:
    model = model_ref_for_angle(class_info, sku_dir, "front") or model_ref_for_angle(class_info, sku_dir, "side") or model_ref_for(class_info, sku_dir)
    return [model, flat_for(class_info, sku_dir, "front"), flat_for(class_info, sku_dir, "back")]


def eagle_base_for(class_info: dict[str, Any], garment: str = "", fit_notes: str = "") -> Path:
    selection = class_info.get("eagle_selection")
    if not isinstance(selection, dict):
        raise ValueError("No model images; complete eagle_selection after visually choosing an Eagle base")
    required_true = ["full_body", "gender_match", "styling_compatible", "batch_unique", "manual_verified"]
    missing = [key for key in required_true if selection.get(key) is not True]
    background = str(selection.get("background") or "").strip().lower()
    if background not in {"pure_color", "seamless_wall", "纯色背景", "无影墙"}:
        missing.append("background=pure_color|seamless_wall")
    if not str(selection.get("uniqueness_signature") or "").strip():
        missing.append("uniqueness_signature")
    if garment == "pants" and not str(selection.get("pants_waistband_preference") or "").strip():
        missing.append("pants_waistband_preference")
    if garment in {"top", "hoodie"} and "crop" not in fit_notes.lower() and selection.get("top_hem_can_hang_naturally") is not True:
        missing.append("top_hem_can_hang_naturally")
    path = str(selection.get("path") or class_info.get("eagle_base") or "").strip()
    if not path:
        missing.append("path")
    if missing:
        raise ValueError(f"Eagle selection is not fully verified: {', '.join(missing)}")
    base = Path(path)
    if not base.exists():
        raise FileNotFoundError(f"Eagle base not found: {base}")
    return base


def model_entries(class_info: dict[str, Any], sku_dir: Path) -> list[tuple[str, Path]]:
    view_map = {"front": "front", "side": "side", "side_3_4": "side", "back": "back"}
    entries = []
    seen = set()
    for img in class_info.get("images", []):
        if img.get("kind") != "model" or not img.get("file"):
            continue
        path = path_by_name(sku_dir, img["file"])
        if path in seen:
            continue
        seen.add(path)
        entries.append((view_map.get(str(img.get("view") or ""), "unknown"), path))
    return entries


def is_recolor_task(text: str) -> bool:
    return ("同款" in text and "换" in text) or any(term in text for term in ["改色", "换色"])


def is_product_edit_task(text: str) -> bool:
    return is_recolor_task(text) or any(term in text for term in ["服装修改", "换装", "产品上身", "服装替换"])


def is_same_person_task(text: str) -> bool:
    return any(term in text for term in ["换成一个人", "换成同一个人", "替换为同一个模特"])


def custom_steps_for(class_info: dict[str, Any], sku_dir: Path) -> list[dict[str, Any]]:
    planned = class_info.get("custom_steps") or []
    if not isinstance(planned, list):
        raise ValueError("custom_steps must be a list")
    steps = []
    known_outputs = set()
    for idx, item in enumerate(planned, 1):
        if not isinstance(item, dict):
            raise ValueError(f"custom_steps[{idx}] must be an object")
        prompt = str(item.get("prompt") or "").strip()
        if not prompt:
            raise ValueError(f"custom_steps[{idx}] is missing a concise prompt")
        refs = []
        for name in item.get("refs", []):
            candidate = Path(str(name))
            refs.append(candidate if candidate.is_absolute() else path_by_name(sku_dir, str(name)))
        refs_from = str(item.get("refs_from") or "").strip()
        if refs_from and refs_from not in known_outputs:
            raise ValueError(f"custom_steps[{idx}] references unavailable prior output: {refs_from}")
        if not refs and not refs_from:
            raise ValueError(f"custom_steps[{idx}] needs refs or refs_from")
        step = {
            "name": str(item.get("name") or f"自定义修改{idx}"),
            "kind": str(item.get("kind") or "custom"),
            "refs": refs,
            "prompt": prompt,
            "intermediate": bool(item.get("intermediate", idx < len(planned))),
        }
        for key in ("refs_from", "flat_angle", "extra_flat_angles", "sets", "move_color_refs"):
            if item.get(key):
                step[key] = item[key]
        output_key = str(step.get("sets") or "").strip()
        if output_key:
            known_outputs.add(output_key)
        steps.append(step)
    return steps


def validate_product_first_dependencies(steps: list[dict[str, Any]]) -> None:
    product_kinds = {"color", "try_on", "product_edit", "garment_replace"}
    transform_kinds = {"pose", "angle_front", "angle_side", "angle_back"}
    has_product_work = any(step.get("kind") in product_kinds for step in steps)
    accurate_outputs: dict[str, bool] = {}
    product_started = False
    for idx, step in enumerate(steps, 1):
        kind = step.get("kind")
        source_key = str(step.get("refs_from") or "")
        source_is_accurate = accurate_outputs.get(source_key, False)
        if kind in product_kinds:
            product_started = True
            output_is_accurate = True
        else:
            output_is_accurate = source_is_accurate
        if has_product_work and kind in transform_kinds and not source_is_accurate:
            raise ValueError(f"Step {idx} ({kind}) must reference a product-accurate output through refs_from")
        output_key = str(step.get("sets") or "")
        if output_key:
            accurate_outputs[output_key] = output_is_accurate
    if has_product_work and not product_started:
        raise ValueError("Product work was requested but no product-correction step was planned")


def build_steps(task: dict[str, Any], class_info: dict[str, Any], image_model: str = "") -> list[dict[str, Any]]:
    sku = task["sku"]
    sku_dir = Path(task["sku_dir"])
    if class_info.get("manual_verified") is not True:
        raise ValueError("Visually inspect the SKU and set manual_verified=true before planning generation")

    custom_steps = custom_steps_for(class_info, sku_dir)
    if custom_steps:
        validate_product_first_dependencies(custom_steps)
        for step in custom_steps:
            validate_generation_prompt(step["prompt"], sku)
        return custom_steps

    items = task.get("actionable_mods") or [task["first_mod"]]
    product_items = [item for item in items if is_product_edit_task(item)]
    pose_items = [item for item in items if "姿势" in item or "动作" in item]
    same_person_items = [item for item in items if is_same_person_task(item)]
    requested_angles = []
    for item in items:
        if any(term in item for term in ["缺失", "缺少", "补充", "增加", "扩充"]):
            for angle in missing_angles(item):
                if angle not in requested_angles:
                    requested_angles.append(angle)

    garment = garment_word(class_info)
    entries = model_entries(class_info, sku_dir)
    steps = []

    if not entries:
        if same_person_items:
            raise ValueError("Same-person replacement needs existing model images")
        if not requested_angles and not pose_items:
            raise ValueError("No model images and no requested view/pose could be planned")
        eagle = eagle_base_for(class_info, garment, fit_notes_for(class_info, "front"))
        steps.append({
            "name": "正面产品上身基准图" if "front" not in requested_angles else "缺失模特正面图",
            "kind": "try_on",
            "refs": [eagle, flat_for(class_info, sku_dir, "front")],
            "prompt": prompt_for("try_on", garment, fit_notes_for(class_info, "front"), image_model),
            "sets": "accurate_front",
            "intermediate": "front" not in requested_angles,
            "source_modification": "；".join(items),
        })
        if "side" in requested_angles:
            steps.append({"name": "缺失模特侧面图", "kind": "angle_side", "refs_from": "accurate_front", "prompt": prompt_for("angle_side", garment, image_model=image_model), "sets": "accurate_side", "source_modification": "；".join(items)})
        if "back" in requested_angles:
            steps.append({"name": "缺失模特背面图", "kind": "angle_back", "refs_from": "accurate_front", "extra_flat_angles": ["front", "back"], "prompt": prompt_for("angle_back", garment, image_model=image_model), "sets": "accurate_back", "source_modification": "；".join(items)})
        for idx, item in enumerate(pose_items, 1):
            steps.append({"name": f"模特姿势更改{idx}", "kind": "pose", "refs_from": "accurate_front", "prompt": prompt_for("pose", garment, image_model=image_model), "sets": f"pose_{idx}", "source_modification": item})
    else:
        corrected_by_path: dict[Path, str] = {}
        corrected_by_angle: dict[str, str] = {}
        product_output_requested = any(is_product_edit_task(item) and not missing_angles(item) and "姿势" not in item and "动作" not in item for item in items)
        if product_items or class_info.get("process_existing_angles_with_product"):
            correction_kind = "color" if any(is_recolor_task(item) for item in product_items) else "try_on"
            for idx, (angle, ref) in enumerate(entries, 1):
                flat_angle = "back" if angle == "back" else "front"
                output_key = f"accurate_{angle}_{idx}"
                step = {
                    "name": f"模特{ {'front': '正面', 'side': '侧面', 'back': '背面'}.get(angle, '') }图产品处理",
                    "kind": correction_kind,
                    "refs": [ref, flat_for(class_info, sku_dir, flat_angle)],
                    "prompt": prompt_for(correction_kind, garment, fit_notes_for(class_info, flat_angle), image_model),
                    "sets": output_key,
                    "intermediate": bool(requested_angles and angle not in requested_angles and not product_output_requested),
                    "source_modification": "；".join(product_items),
                }
                if correction_kind == "color":
                    step["move_color_refs"] = [str(ref)]
                steps.append(step)
                corrected_by_path[ref] = output_key
                corrected_by_angle.setdefault(angle, output_key)

        base_angle, base_ref = next(((angle, ref) for angle, ref in entries if angle == "front"), entries[0])
        base_output = corrected_by_path.get(base_ref)

        for idx, item in enumerate(same_person_items, 1):
            base_name = class_info.get("same_person_base")
            targets = class_info.get("same_person_targets") or []
            if not base_name or not targets:
                raise ValueError("Set same_person_base and same_person_targets after visually choosing the identity base")
            identity_ref = path_by_name(sku_dir, base_name)
            for target_idx, target_name in enumerate(targets, 1):
                target_ref = path_by_name(sku_dir, target_name)
                step = {"name": f"三张模特图换成同一个人{target_idx}", "kind": "same_person", "refs": [target_ref, identity_ref], "prompt": prompt_for("same_person", garment, image_model=image_model), "sets": f"same_person_{idx}_{target_idx}", "source_modification": item}
                if target_ref in corrected_by_path:
                    step["refs_from"] = corrected_by_path[target_ref]
                    step["refs"] = [identity_ref]
                steps.append(step)

        existing_angles = {angle for angle, _ in entries}
        for angle in requested_angles:
            if angle in existing_angles:
                continue
            angle_kind = {"front": "angle_front", "side": "angle_side", "back": "angle_back"}[angle]
            view_name = {"front": "正面", "side": "侧面", "back": "背面"}[angle]
            output_key = f"accurate_{angle}"
            step = {"name": f"缺失模特{view_name}图", "kind": angle_kind, "prompt": prompt_for(angle_kind, garment, image_model=image_model), "sets": output_key, "source_modification": "；".join(items)}
            if base_output:
                step["refs_from"] = base_output
            else:
                step["refs"] = [base_ref]
            if angle == "back":
                step["extra_flat_angles"] = ["front", "back"]
            steps.append(step)

        pose_targets = class_info.get("pose_targets") or [str(base_ref.name)]
        for idx, item in enumerate(pose_items, 1):
            for target_idx, target_name in enumerate(pose_targets, 1):
                target_ref = path_by_name(sku_dir, target_name)
                step = {"name": f"模特姿势更改{idx}_{target_idx}", "kind": "pose", "prompt": prompt_for("pose", garment, image_model=image_model), "sets": f"pose_{idx}_{target_idx}", "source_modification": item}
                if target_ref in corrected_by_path:
                    step["refs_from"] = corrected_by_path[target_ref]
                elif base_output and target_ref == base_ref:
                    step["refs_from"] = base_output
                else:
                    step["refs"] = [target_ref]
                steps.append(step)

    if not steps:
        raise ValueError("No executable steps were produced; add visually planned custom_steps")
    validate_product_first_dependencies(steps)
    for step in steps:
        validate_generation_prompt(step["prompt"], sku)
    return steps


def collect_result_urls(payload: Any) -> list[str]:
    urls = []
    if isinstance(payload, dict):
        for key in ("url", "image_url", "output_url"):
            if isinstance(payload.get(key), str):
                urls.append(payload[key])
        for key in ("data", "images", "output", "result", "results"):
            value = payload.get(key)
            if isinstance(value, list):
                for item in value:
                    urls.extend(collect_result_urls(item))
            elif isinstance(value, dict):
                urls.extend(collect_result_urls(value))
            elif isinstance(value, str) and value.startswith("http"):
                urls.append(value)
    return list(dict.fromkeys(urls))


def collect_b64(payload: Any) -> list[str]:
    out = []
    if isinstance(payload, dict):
        for key in ("b64_json", "base64", "image_base64"):
            value = payload.get(key)
            if isinstance(value, str) and len(value) > 100:
                out.append(value)
        for value in payload.values():
            if isinstance(value, (dict, list)):
                out.extend(collect_b64(value))
    elif isinstance(payload, list):
        for item in payload:
            out.extend(collect_b64(item))
    return out


def poll_task(task_id: str, api_key: str) -> Any:
    deadline = time.time() + REQUEST_TIMEOUT
    last = None
    while time.time() < deadline:
        resp = request_with_retry("GET", f"{API_BASE}/images/generations/{task_id}", headers={"Authorization": f"Bearer {api_key}"})
        payload = safe_json(resp)
        last = payload
        status = str(payload.get("status", "")).lower() if isinstance(payload, dict) else ""
        if status in {"succeeded", "success", "completed", "complete", "failed", "error", "cancelled"}:
            return payload
        if collect_result_urls(payload) or collect_b64(payload):
            return payload
        time.sleep(POLL_INTERVAL)
    raise TimeoutError(f"generation task {task_id} timed out; last={last}")


def generate_image(prompt: str, refs: list[Path], api_key: str, model: str) -> dict[str, Any]:
    uploads = [upload_image(ref, api_key) for ref in refs]
    body = {"model": model, "prompt": prompt, "size": SIZE, "n": 1, "image_urls": [u["url"] for u in uploads], "metadata": {"resolution": RESOLUTION}}
    if model == "gpt-image-2-vip":
        body = {"model": model, "prompt": prompt, "n": 1, "size": SIZE, "resolution": RESOLUTION, "quality": "medium", "image_urls": [u["url"] for u in uploads], "response_format": "url"}
    resp = request_with_retry("POST", f"{API_BASE}/images/generations", headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"}, json=body)
    payload = safe_json(resp)
    if resp.status_code >= 400:
        raise RuntimeError(f"generation failed {resp.status_code}: {payload}")
    task_id = payload.get("id") if isinstance(payload, dict) else None
    status = str(payload.get("status", "")).lower() if isinstance(payload, dict) else ""
    result = payload
    if task_id and not collect_result_urls(payload) and status not in {"succeeded", "success", "completed", "complete"}:
        result = poll_task(task_id, api_key)
    return {"uploads": uploads, "request": body, "initial": payload, "result": result}


def save_outputs(result: dict[str, Any], out_base: Path) -> list[str]:
    payload = result.get("result", result)
    urls = collect_result_urls(payload)
    saved = []
    for idx, url in enumerate(urls, 1):
        ext = ".jpg"
        m = re.search(r"\.(jpg|jpeg|png|webp)(?:\?|$)", url, re.I)
        if m:
            ext = "." + m.group(1).lower().replace("jpeg", "jpg")
        out = out_base.with_name(out_base.name + (f"_{idx}" if len(urls) > 1 else "") + ext)
        resp = request_with_retry("GET", url)
        if resp.status_code >= 400:
            raise RuntimeError(f"download failed {resp.status_code}: {url}")
        out.write_bytes(resp.content)
        saved.append(str(out))
    for idx, b64 in enumerate(collect_b64(payload), 1):
        if b64.strip().startswith("data:") and "," in b64:
            b64 = b64.split(",", 1)[1]
        out = out_base.with_name(out_base.name + (f"_{idx}" if len(saved) > 0 else "") + ".png")
        out.write_bytes(base64.b64decode(b64))
        saved.append(str(out))
    return saved


def archive_old(root: Path, sku_names: set[str] | None = None) -> Path:
    archive = root / f"旧AIGC生成结果_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
    archive.mkdir(exist_ok=True)
    for sku_dir in [p for p in root.iterdir() if p.is_dir() and not p.name.startswith("旧AIGC生成结果_") and p.name != "AIGC视觉识别拼版"]:
        if sku_names is not None and sku_dir.name not in sku_names:
            continue
        files = list(sku_dir.glob("AIGC+*"))
        if not files:
            continue
        dest = archive / sku_dir.name
        dest.mkdir(exist_ok=True)
        for f in files:
            shutil.move(str(f), str(dest / f.name))
    return archive


def move_to_process_dir(path: Path, process_dir: Path) -> Path:
    if not path.exists():
        return path
    dest = process_dir / path.name
    if dest.exists():
        stem, suffix = path.stem, path.suffix
        idx = 1
        while True:
            candidate = process_dir / f"{stem}_原参考{idx}{suffix}"
            if not candidate.exists():
                dest = candidate
                break
            idx += 1
    shutil.move(str(path), str(dest))
    return dest


def eagle_duplicate_errors(tasks: list[dict[str, Any]], classes: dict[str, Any]) -> dict[str, str]:
    seen: dict[str, str] = {}
    errors: dict[str, str] = {}
    for task in tasks:
        sku = task["sku"]
        class_info = classes.get(sku) or {}
        if any(img.get("kind") == "model" for img in class_info.get("images", [])):
            continue
        selection = class_info.get("eagle_selection") or {}
        path = str(selection.get("path") or class_info.get("eagle_base") or "").strip().lower()
        signature = str(selection.get("uniqueness_signature") or "").strip().lower()
        for label, value in (("path", path), ("visual signature", signature)):
            if not value:
                continue
            key = f"{label}:{value}"
            if key in seen and seen[key] != sku:
                message = f"Duplicate Eagle base {label} with SKU {seen[key]}"
                errors[sku] = message
                errors.setdefault(seen[key], f"Duplicate Eagle base {label} with SKU {sku}")
            else:
                seen[key] = sku
    return errors


def process_sku(task: dict[str, Any], class_info: dict[str, Any], api_key: str, model: str, dry_run: bool) -> dict[str, Any]:
    sku_dir = Path(task["sku_dir"])
    process_dir = sku_dir / "过程文件"
    process_dir.mkdir(exist_ok=True)
    log_path = sku_dir / f"AIGC处理日志_{task['sku']}.txt"
    record = {
        "sku": task["sku"],
        "first_mod": task["first_mod"],
        "actionable_mods": task.get("actionable_mods", [task["first_mod"]]),
        "skipped_mods": task.get("skipped_mods", []),
        "status": "started",
        "outputs": [],
        "errors": [],
    }
    produced: dict[str, Path] = {}
    color_refs_to_move: set[Path] = set()
    try:
        steps = build_steps(task, class_info, model)
    except Exception as exc:
        record["status"] = "failed"
        record["errors"].append(repr(exc))
        append_log(log_path, {"time": now(), "event": "plan_error", "error": repr(exc)})
        return record
    append_log(log_path, {"time": now(), "event": "sku_start", "task": task, "steps": [{"name": s["name"], "prompt": s["prompt"], "refs": [str(r) for r in s.get("refs", [])], "refs_from": s.get("refs_from"), "flat_angle": s.get("flat_angle"), "extra_flat_angles": s.get("extra_flat_angles", [])} for s in steps]})
    if dry_run:
        record["status"] = "planned"
        return record
    for idx, step in enumerate(steps, 1):
        try:
            refs = list(step.get("refs", []))
            if "refs_from" in step:
                refs = [produced[step["refs_from"]], *refs]
                for extra_angle in step.get("extra_flat_angles", []):
                    refs.append(flat_for(class_info, sku_dir, extra_angle))
                if step.get("flat_angle"):
                    refs.append(flat_for(class_info, sku_dir, step["flat_angle"]))
            out_parent = process_dir if step.get("intermediate") else sku_dir
            out_base = out_parent / f"AIGC+{sanitize_name(step['name'])}+{idx}"
            append_log(log_path, {"time": now(), "event": "subtask_start", "step": step, "output_base": str(out_base), "prompt": step["prompt"], "reference_images": [str(r) for r in refs]})
            result = generate_image(step["prompt"], refs, api_key, model)
            saved = save_outputs(result, out_base)
            append_log(log_path, {"time": now(), "event": "generation_result", "result": result, "saved_outputs": saved})
            if not saved:
                raise RuntimeError("generation completed but no output image was found")
            record["outputs"].extend(saved)
            if step.get("sets"):
                produced[step["sets"]] = Path(saved[0])
            if step.get("kind") == "color" and refs:
                source_ref = Path(refs[0])
                if source_ref.parent == sku_dir and not source_ref.name.startswith("AIGC+"):
                    color_refs_to_move.add(source_ref)
            for ref in step.get("move_color_refs", []):
                ref_path = Path(ref)
                if ref_path.parent == sku_dir and ref_path.name.startswith("AIGC+") is False:
                    color_refs_to_move.add(ref_path)
        except Exception as exc:
            record["errors"].append({"step": step.get("name"), "error": repr(exc)})
            append_log(log_path, {"time": now(), "event": "subtask_error", "step": step, "error": repr(exc)})
    moved_refs = []
    if not record["errors"]:
        for ref_path in sorted(color_refs_to_move, key=lambda p: p.name):
            try:
                moved = move_to_process_dir(ref_path, process_dir)
                moved_refs.append({"from": str(ref_path), "to": str(moved)})
            except Exception as exc:
                record["errors"].append({"step": "move_color_reference", "path": str(ref_path), "error": repr(exc)})
                append_log(log_path, {"time": now(), "event": "move_color_reference_error", "path": str(ref_path), "error": repr(exc)})
        if moved_refs:
            append_log(log_path, {"time": now(), "event": "moved_color_reference_images", "moved": moved_refs})
    record["status"] = "success" if record["outputs"] and not record["errors"] else ("partial" if record["outputs"] else "failed")
    return record


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--root", required=True)
    ap.add_argument("--excel", default="新建 Microsoft Excel 工作表.xlsx")
    ap.add_argument("--mode", choices=["classify", "plan", "generate"], required=True)
    ap.add_argument("--api-key", default=os.environ.get("TOAPIS_API_KEY", ""))
    ap.add_argument("--image-model", default=DEFAULT_IMAGE_MODEL)
    ap.add_argument("--vision-model", default=DEFAULT_VISION_MODEL)
    ap.add_argument("--workers", type=int, default=4)
    ap.add_argument("--archive-old", action="store_true")
    args = ap.parse_args()

    root = Path(args.root)
    tasks = read_tasks(root, args.excel)
    print(json.dumps({"task_count": len(tasks), "skus": [t["sku"] for t in tasks]}, ensure_ascii=False), flush=True)

    if args.mode == "classify":
        out = classify(root, tasks, args.api_key or None, args.vision_model)
        print(json.dumps({"classification": str(out), "contact_sheets": str(root / "AIGC视觉识别拼版")}, ensure_ascii=False), flush=True)
        return 0

    classes = load_classification(root)
    duplicate_eagle = eagle_duplicate_errors(tasks, classes)
    if args.archive_old and args.mode == "generate":
        archive = archive_old(root, {t["sku"] for t in tasks})
        print(json.dumps({"archive_old": str(archive)}, ensure_ascii=False), flush=True)
    if args.mode == "generate" and not args.api_key:
        raise SystemExit("TOAPIS_API_KEY is required for generate mode")

    summary = {"started_at": now(), "root": str(root), "total_skus": len(tasks), "skus": []}
    dry_run = args.mode == "plan"
    with ThreadPoolExecutor(max_workers=args.workers) as ex:
        futures = {}
        for task in tasks:
            ci = classes.get(task["sku"])
            if not ci:
                summary["skus"].append({"sku": task["sku"], "status": "failed", "errors": ["missing classification"]})
                continue
            if task["sku"] in duplicate_eagle:
                summary["skus"].append({"sku": task["sku"], "status": "failed", "errors": [duplicate_eagle[task["sku"]]]})
                continue
            futures[ex.submit(process_sku, task, ci, args.api_key, args.image_model, dry_run)] = task
        for fut in as_completed(futures):
            rec = fut.result()
            summary["skus"].append(rec)
            print(json.dumps(rec, ensure_ascii=False), flush=True)
    summary["finished_at"] = now()
    out = root / ("AIGC任务计划.json" if dry_run else "AIGC批处理汇总.json")
    out.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"summary": str(out)}, ensure_ascii=False), flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
